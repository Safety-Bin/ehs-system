"""근거를 남기는 삼값 조건 평가와 검토용 파일 출력."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
import json
import math

import pandas as pd

from profiles import FLAG_FIELDS, NUMERIC_FIELDS, PROCESS_OPTIONS, TEXT_FIELDS, normalize_profile


def is_unknown(value):
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() in ("", "미확인", "미입력", "알 수 없음")
    if isinstance(value, (list, tuple, dict)):
        return len(value) == 0
    if isinstance(value, float):
        return not math.isfinite(value)
    return type(value).__name__ in ("NAType", "NaTType")


def field_label(field):
    return (NUMERIC_FIELDS.get(field) or FLAG_FIELDS.get(field) or TEXT_FIELDS.get(field) or {}).get("label", field)


def _lookup(profile, path):
    value = profile
    for part in path.split("."):
        if not isinstance(value, dict):
            return None
        value = value.get(part)
    return value


def evaluate_condition(condition, profile):
    """Return (True/False/None, evidence, missing_fields) using Kleene logic."""
    if not condition:
        return None, [], ["적용 조건"]
    for operator in ("all", "any"):
        if operator in condition:
            children = condition[operator]
            if not children:
                return None, [], ["적용 조건"]
            evaluated = [evaluate_condition(child, profile) for child in children]
            values = [item[0] for item in evaluated]
            if operator == "all":
                value = False if False in values else (None if None in values else True)
            else:
                value = True if True in values else (None if None in values else False)
            evidence = list(dict.fromkeys(text for result in evaluated for text in result[1]))
            missing = list(dict.fromkeys(text for result in evaluated for text in result[2]))
            return value, evidence, missing
    if "not" in condition:
        value, evidence, missing = evaluate_condition(condition["not"], profile)
        return (None if value is None else not value), evidence, missing
    field = condition.get("field", "")
    value = _lookup(profile, field)
    if is_unknown(value):
        return None, [], [field]
    operation = condition.get("op", "eq")
    target = condition.get("value")
    try:
        if operation == "eq":
            result = value == target
        elif operation == "ne":
            result = value != target
        elif operation in ("gte", "gt", "lte", "lt"):
            if isinstance(value, bool) or isinstance(target, bool):
                return None, [], [field]
            a, b = float(value), float(target)
            if not math.isfinite(a) or not math.isfinite(b):
                return None, [], [field]
            result = {"gte": a >= b, "gt": a > b, "lte": a <= b, "lt": a < b}[operation]
        elif operation == "contains":
            result = target in value
        elif operation == "intersects":
            result = bool(set(value).intersection(target))
        elif operation == "truthy":
            result = bool(value)
        else:
            return None, [], [field]
    except (ValueError, TypeError, KeyError):
        return None, [], [field]
    shown = "예" if value is True else "아니오" if value is False else str(value)
    if field == "processes":
        shown = ", ".join(PROCESS_OPTIONS.get(item, item) for item in value)
    return bool(result), [f"{field_label(field)}: {shown}"], []


def condition_fields(condition):
    if not condition:
        return set()
    if "field" in condition:
        return {condition["field"]}
    if "not" in condition:
        return condition_fields(condition["not"])
    return set().union(*(condition_fields(child) for op in ("all", "any") for child in condition.get(op, [])))


def build_checklist(profile, rules=None):
    if rules is None:
        from legal_catalogue import RULES
        rules = RULES
    normalized = normalize_profile(profile)
    rows = []
    for rule in rules:
        value, matched, missing = evaluate_condition(rule.get("condition"), normalized)
        if value is False:
            status = "현재정보상비대상"
        elif value is None or rule.get("force_review", False):
            status = "추가확인"
        else:
            status = "검토대상"
        scope_note = rule.get("scope_note", "")
        if rule.get("force_review") and value is not False:
            scope_note = "대상 분류·예외·세부 제원 확인 필요. " + scope_note
        rows.append({
            "rule_id": rule.get("id", rule.get("rule_id", "")),
            "category": rule.get("category", "기타"),
            "title": rule.get("title", "제목 없음"),
            "applicability": status,
            "law_name": rule.get("law_name", ""),
            "article": rule.get("article", ""),
            "source_url": rule.get("source_url", ""),
            "scope_note": scope_note,
            "actions": list(rule.get("actions", [])),
            "evidence": list(rule.get("evidence", [])),
            "rule_kind": rule.get("rule_kind", "실무 확인"),
            "matched_fields": matched,
            "missing_fields": [field_label(item) for item in missing],
            "source_checked_on": rule.get("source_checked_on", "미확인"),
        })
    return rows


def safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, (list, dict, tuple)):
        value = json.dumps(value, ensure_ascii=False, default=str) if isinstance(value, dict) else "\n".join(str(item) for item in value)
    if isinstance(value, str):
        value = value[:32000]
        if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
            return "'" + value
    return value


def checklist_dataframe(rows, dispositions=None):
    dispositions = dispositions or {}
    records = []
    for row in rows:
        disposition = dispositions.get(row["rule_id"], {})
        records.append({
            "ID": row["rule_id"], "분야": row["category"], "검토 구분": row["rule_kind"],
            "관련성": row["applicability"], "확인 항목": row["title"],
            "입력 근거": "\n".join(row["matched_fields"]),
            "추가 입력": "\n".join(row["missing_fields"]),
            "적용 조건·예외": row["scope_note"],
            "확인할 행동": "\n".join(row["actions"]), "준비할 자료": "\n".join(row["evidence"]),
            "법령": row["law_name"], "조문": row["article"], "공식 출처": row["source_url"],
            "출처 확인일": row["source_checked_on"],
            "진행 상태": disposition.get("status", "미착수"),
            "담당자": disposition.get("owner", ""), "기한": disposition.get("due_date", ""),
            "검토 메모": disposition.get("memo", ""),
        })
    return pd.DataFrame([{key: safe_cell(value) for key, value in item.items()} for item in records])


def _table_dataframe(records):
    return pd.DataFrame([{str(key): safe_cell(value) for key, value in item.items()} for item in records])


def export_workbook(profile, rows, impact=None, dispositions=None):
    output = BytesIO()
    scalar_rows = []
    for key, value in profile.items():
        if key in ("buildings", "chemicals", "equipment"):
            continue
        scalar_rows.append({"항목": field_label(key), "값": "미확인" if value is None else safe_cell(value)})
    notices = [
        {"구분": "출력일시", "내용": datetime.now(timezone(timedelta(hours=9))).isoformat()},
        {"구분": "판단 범위", "내용": "입력 정보와 등록된 규칙의 관련성을 선별한 검토 자료. 이행 완료·법적 면제의 확정 결과가 아님."},
        {"구분": "법령 상태", "내용": "출처 확인일 이후 개정·부칙·경과조치와 행정해석은 공식 원문에서 확인."},
        {"구분": "알 수 없는 입력", "내용": "미입력·미확인을 아니오/0/영향 없음으로 해석하지 않음."},
        {"구분": "표준 범위", "내용": "작업 안전·보건, 도급·건설, 건축·소방, 위험물·화학물질·환경의 등록 항목."},
    ]
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _table_dataframe(notices).to_excel(writer, index=False, sheet_name="검토 안내")
        _table_dataframe(scalar_rows).to_excel(writer, index=False, sheet_name="사업장 제원")
        checklist_dataframe(rows, dispositions).to_excel(writer, index=False, sheet_name="체크리스트")
        for key, label in (("buildings", "건축물 상세"), ("chemicals", "물질 상세"), ("equipment", "설비 상세")):
            table = _table_dataframe(profile.get(key, []))
            if table.empty:
                table = pd.DataFrame([{"입력 상태": "상세 정보 미입력"}])
            table.to_excel(writer, index=False, sheet_name=label)
        if impact:
            _table_dataframe([{"구분": "요약", "내용": impact.get("summary", "")},
                              {"구분": "상태", "내용": impact.get("status", "")},
                              {"구분": "확인사항", "내용": impact.get("warnings", [])}]).to_excel(writer, index=False, sheet_name="개정 검토 요약")
            results = impact.get("results", [])
            if results:
                _table_dataframe(results).to_excel(writer, index=False, sheet_name="개정 영향성")
        from openpyxl.styles import Alignment, Font, PatternFill
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="163D50")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[1].height = 30
            for cells in sheet.iter_cols():
                heading = str(cells[0].value or "")
                width = 18 if len(heading) < 6 else 26
                if heading in ("확인 항목", "적용 조건·예외", "확인할 행동", "준비할 자료", "내용", "quote", "actions", "입력 근거", "검토 메모"):
                    width = 46
                if heading in ("공식 출처", "source_url"):
                    width = 54
                sheet.column_dimensions[cells[0].column_letter].width = width
                for cell in cells[1:]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for i in range(2, sheet.max_row + 1):
                sheet.row_dimensions[i].height = 58
    return output.getvalue()
