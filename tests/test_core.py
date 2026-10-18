import sys
from pathlib import Path
import unittest
from io import BytesIO

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from core import build_checklist, evaluate_condition, export_workbook
from profiles import get_template, normalize_profile


class TriStateTests(unittest.TestCase):
    def test_zero_is_known_and_empty_is_unknown(self):
        clause = {"field": "headcount", "op": "gte", "value": 5}
        self.assertIs(evaluate_condition(clause, {"headcount": 0})[0], False)
        self.assertIsNone(evaluate_condition(clause, {"headcount": None})[0])

    def test_false_hazard_is_not_missing(self):
        clause = {"field": "confined_space", "op": "eq", "value": True}
        self.assertIs(evaluate_condition(clause, {"confined_space": False})[0], False)
        self.assertIsNone(evaluate_condition(clause, {"confined_space": None})[0])

    def test_logical_short_circuit_preserves_unknown(self):
        a = {"field": "confined_space", "op": "eq", "value": True}
        b = {"field": "headcount", "op": "gte", "value": 5}
        self.assertIs(evaluate_condition({"all": [a, b]}, {"confined_space": False})[0], False)
        self.assertIsNone(evaluate_condition({"any": [a, b]}, {"confined_space": False})[0])

    def test_numeric_boundary(self):
        clause = {"field": "headcount", "op": "gte", "value": 5}
        self.assertIs(evaluate_condition(clause, {"headcount": 5})[0], True)

    def test_force_review_does_not_create_legal_verdict(self):
        rule = {"id": "r", "title": "복합 대상", "force_review": True,
                "condition": {"field": "headcount", "op": "gte", "value": 5}}
        self.assertEqual(build_checklist({"headcount": 5}, [rule])[0]["applicability"], "추가확인")

    def test_unknown_profile_cannot_be_marked_exempt(self):
        rule = {"id": "r", "condition": {"field": "psm_confirmed", "op": "eq", "value": True}}
        self.assertEqual(build_checklist({}, [rule])[0]["applicability"], "추가확인")

    def test_profile_distinguishes_imported_false_and_missing(self):
        profile = normalize_profile({"confined_space": "아니오", "headcount": 0,
                                     "ground_floors": float("nan")})
        self.assertIs(profile["confined_space"], False)
        self.assertEqual(profile["headcount"], 0)
        self.assertIsNone(profile["ground_floors"])

    def test_workbook_does_not_execute_user_formulas(self):
        from openpyxl import load_workbook
        profile = normalize_profile({"name": "=1+1"})
        rule = {"id": "r", "title": "자료 확인", "condition": {"field": "headcount", "op": "gte", "value": 1}}
        rows = build_checklist(profile, [rule])
        raw = export_workbook(profile, rows, dispositions={"r": {"memo": "=HYPERLINK(\"x\")"}})
        book = load_workbook(BytesIO(raw), data_only=False)
        for sheet in book:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")
        self.assertIn("체크리스트", book.sheetnames)

    def test_templates_have_unique_identity_and_unknown_specs(self):
        a, b = get_template("pcb"), get_template("pcb")
        self.assertNotEqual(a["uid"], b["uid"])
        self.assertIsNone(a["headcount"])
        self.assertIsNone(a["psm_confirmed"])


if __name__ == "__main__":
    unittest.main()
